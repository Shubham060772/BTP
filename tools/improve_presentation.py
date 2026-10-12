from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


SOURCE = Path("BTP_AdaFuse_Presentation.pptx")
OUTPUT = Path("BTP_AdaFuse_Presentation_Improved.pptx")


def rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor.from_string(value.upper())


BG = rgb("0A1220")
CARD = rgb("111E2F")
CARD_ALT = rgb("15243A")
BORDER = rgb("263A52")
TEXT = rgb("E7EEF8")
MUTED = rgb("95A8BE")
BLUE = rgb("38BDF8")
TEAL = rgb("2DD4BF")
ORANGE = rgb("F97316")
GREEN = rgb("22C55E")
PURPLE = rgb("A78BFA")
YELLOW = rgb("FBBF24")
PINK = rgb("F472B6")
RED = rgb("FB7185")


COLOR_MAP = {
    "CCD6E8": TEXT,
    "FFFFFF": TEXT,
    "8899AA": MUTED,
    "1E90FF": BLUE,
    "00D4AA": TEAL,
    "FF6B35": ORANGE,
    "2ECC71": GREEN,
    "AA44FF": PURPLE,
    "FFAA00": YELLOW,
    "FFCC00": YELLOW,
    "FF4477": PINK,
    "44FF99": GREEN,
}

FILL_MAP = {
    "16273E": CARD_ALT,
    "122032": CARD,
    "1E90FF": BLUE,
    "00D4AA": TEAL,
    "FF6B35": ORANGE,
    "2ECC71": GREEN,
    "AA44FF": PURPLE,
    "8899AA": MUTED,
    "FFAA00": YELLOW,
    "FFCC00": YELLOW,
    "FF4477": PINK,
    "44FF99": GREEN,
    "0A201A": rgb("0D2A23"),
    "0A200F": rgb("0D2819"),
    "1E3020": rgb("163223"),
    "261A0A": rgb("302315"),
    "26100A": rgb("321815"),
    "0A1426": rgb("101D32"),
}


def shape_text(shape) -> str:
    return shape.text.strip() if getattr(shape, "has_text_frame", False) else ""


def set_box(shape, x=None, y=None, w=None, h=None):
    if x is not None:
        shape.left = Inches(x)
    if y is not None:
        shape.top = Inches(y)
    if w is not None:
        shape.width = Inches(w)
    if h is not None:
        shape.height = Inches(h)


def recolor(shape):
    try:
        if shape.fill.type is not None:
            current = str(shape.fill.fore_color.rgb)
            if current in FILL_MAP:
                shape.fill.solid()
                shape.fill.fore_color.rgb = FILL_MAP[current]
    except (AttributeError, TypeError):
        pass

    if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
        try:
            fill_rgb = str(shape.fill.fore_color.rgb)
            if fill_rgb in {str(CARD), str(CARD_ALT)}:
                shape.line.color.rgb = BORDER
                shape.line.width = Pt(0.8)
            else:
                shape.line.fill.background()
        except (AttributeError, TypeError):
            pass

    if not getattr(shape, "has_text_frame", False):
        return

    shape.text_frame.word_wrap = True
    for paragraph in shape.text_frame.paragraphs:
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Aptos"
            try:
                current = str(run.font.color.rgb)
                if current in COLOR_MAP:
                    run.font.color.rgb = COLOR_MAP[current]
            except (AttributeError, TypeError):
                pass


def title_style(shape, title: bool):
    if not getattr(shape, "has_text_frame", False):
        return
    shape.text_frame.margin_left = 0
    shape.text_frame.margin_right = 0
    shape.text_frame.margin_top = 0
    shape.text_frame.margin_bottom = 0
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for p in shape.text_frame.paragraphs:
        for run in p.runs:
            run.font.name = "Aptos Display" if title else "Aptos"
            run.font.size = Pt(27 if title else 12.5)
            run.font.bold = title
            run.font.color.rgb = TEXT if title else MUTED


def style_headers(prs):
    for idx, slide in enumerate(prs.slides, 1):
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = BG
        for shape in slide.shapes:
            recolor(shape)

        if idx in (1, 15):
            continue

        candidates = [
            s for s in slide.shapes
            if getattr(s, "has_text_frame", False) and s.top < Inches(1.0) and shape_text(s)
        ]
        candidates.sort(key=lambda s: s.top)
        if candidates:
            title = candidates[0]
            set_box(title, 0.55, 0.20, 11.7, 0.48)
            title_style(title, True)
        if len(candidates) > 1:
            subtitle = candidates[1]
            set_box(subtitle, 0.55, 0.72, 11.7, 0.30)
            title_style(subtitle, False)

        # Make the divider and top accent precise and consistent.
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.LINE and shape.top < Inches(1.2):
                set_box(shape, 0.55, 1.08, 12.23, 0)
                shape.line.color.rgb = BORDER
                shape.line.width = Pt(1)
            elif shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and shape.top < Inches(0.12):
                set_box(shape, 0, 0, 13.333, 0.06)
                shape.width = prs.slide_width
                shape.fill.solid()
                shape.fill.fore_color.rgb = BLUE
                shape.line.fill.background()

        # Compact slide index in the header; avoids adding a footer over dense slides.
        page = slide.shapes.add_textbox(Inches(12.35), Inches(0.24), Inches(0.42), Inches(0.32))
        page.text_frame.clear()
        page.text_frame.margin_left = page.text_frame.margin_right = 0
        page.text_frame.margin_top = page.text_frame.margin_bottom = 0
        page.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = page.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = f"{idx:02d}"
        r.font.name = "Aptos"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = BLUE


def fix_title_slide(slide):
    for shape in slide.shapes:
        recolor(shape)
        text = shape_text(shape)
        if text == "AdaFuse":
            set_box(shape, 1.0, 1.18, 11.33, 1.05)
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name = "Aptos Display"
                    r.font.size = Pt(68)
                    r.font.bold = True
                    r.font.color.rgb = BLUE
        elif text.startswith("Adaptive Attention Fusion"):
            set_box(shape, 1.0, 2.32, 11.33, 0.55)
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(22)
                    r.font.color.rgb = TEXT
        elif text.startswith("BTP —"):
            set_box(shape, 1.0, 3.42, 11.33, 0.42)
        elif text == "Indian Institute of Technology Patna":
            set_box(shape, 1.0, 4.15, 11.33, 0.34)
        elif text.startswith("Department of Computer Science"):
            set_box(shape, 1.0, 4.55, 11.33, 0.34)

    badges = [s for s in slide.shapes if shape_text(s) in {
        "Transformers", "Graph Neural Networks", "Adaptive Fusion"
    }]
    xs = [2.22, 5.37, 8.52]
    for shape, x in zip(badges, xs):
        set_box(shape, x, 5.54, 2.6, 0.42)
        shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        for p in shape.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(11)
                r.font.bold = True
    badge_bgs = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and 5.3 < s.top / 914400 < 6.1]
    for shape, x in zip(badge_bgs, xs):
        set_box(shape, x, 5.53, 2.6, 0.44)


def fix_outline(slide):
    labels = [
        "Problem Motivation & Research Gap",
        "Literature Review — 26+ Papers",
        "AdaFuse Architecture",
        "Experimental Setup",
        "Results & Analysis",
        "What Went Wrong — Honest Assessment",
        "Next Steps & Timeline",
    ]
    positions = [
        (0.55, 1.34), (6.83, 1.34),
        (0.55, 2.73), (6.83, 2.73),
        (0.55, 4.12), (6.83, 4.12),
        (3.69, 5.51),
    ]
    for number, (label, (x, y)) in enumerate(zip(labels, positions), 1):
        label_shape = next(s for s in slide.shapes if shape_text(s) == label)
        num_shape = next(s for s in slide.shapes if shape_text(s) == f"{number:02d}")
        old_x, old_y = num_shape.left, num_shape.top
        backgrounds = [
            s for s in slide.shapes
            if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
            and abs(s.left - old_x) < Inches(0.03)
            and abs(s.top - (old_y - Inches(0.40))) < Inches(0.04)
        ]
        set_box(label_shape, x + 0.90, y + 0.12, 4.78, 0.78)
        label_shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        set_box(num_shape, x, y + 0.29, 0.72, 0.48)
        num_shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        for p in num_shape.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
        for bg in backgrounds:
            if bg.width > Inches(2):
                set_box(bg, x, y, 5.95, 1.08)
            else:
                set_box(bg, x, y, 0.72, 1.08)


def fix_slide_3(slide):
    for shape in slide.shapes:
        text = shape_text(shape)
        if text == "⟷":
            set_box(shape, 6.25, 2.45, 0.83, 0.42)
        elif text.startswith("Complementary"):
            set_box(shape, 5.91, 2.91, 1.52, 0.48)
            shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE


def fix_experimental_setup(slide):
    old_rows = [4.30, 4.78, 5.15, 5.52, 5.89, 6.26, 6.63, 7.00]
    new_rows = [4.14] + [4.56 + 0.35 * i for i in range(7)]
    for shape in slide.shapes:
        y = shape.top / 914400
        nearest = min(range(len(old_rows)), key=lambda i: abs(y - old_rows[i]))
        if abs(y - old_rows[nearest]) < 0.08:
            is_bg = shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
            h = 0.37 if nearest == 0 else 0.32
            set_box(shape, y=new_rows[nearest], h=h if is_bg else h - 0.02)
            if getattr(shape, "has_text_frame", False):
                shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size and r.font.size.pt > 11:
                            r.font.size = Pt(10.5)


def fix_gate_table(slide):
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        x = shape.left / 914400
        y = shape.top / 914400
        if 1.25 <= y <= 4.35 and x < 6.0:
            if 2.35 <= x < 3.55:
                set_box(shape, x=2.48, w=1.12)
            elif 3.55 <= x < 4.75:
                set_box(shape, x=3.65, w=1.05)
            elif x >= 4.75:
                set_box(shape, x=4.77, w=0.98)


def align_dense_cards(slide, slide_num):
    if slide_num not in {5, 7, 11, 12, 13}:
        return
    # Vertically center labels and severity/number chips in repeated horizontal cards.
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.top > Inches(1.15):
            text = shape_text(shape)
            if text in {"1", "2", "3", "4", "5", "CRITICAL", "HIGH", "MEDIUM", "LOW", "✓"}:
                shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                for p in shape.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.CENTER


def fix_summary_cards(slide):
    # Equal inner padding and readable card body copy.
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and 1.8 < shape.top / 914400 < 5.1:
            shape.text_frame.margin_left = Inches(0.03)
            shape.text_frame.margin_right = Inches(0.03)
            shape.text_frame.margin_top = Inches(0.03)
            shape.text_frame.margin_bottom = Inches(0.03)
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size and r.font.size.pt < 12:
                        r.font.size = Pt(12)


def fix_end_slide(slide):
    for shape in slide.shapes:
        recolor(shape)
        text = shape_text(shape)
        if text == "Thank You":
            set_box(shape, 0.8, 1.62, 11.73, 0.95)
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name = "Aptos Display"
                    r.font.size = Pt(58)
                    r.font.bold = True
                    r.font.color.rgb = TEXT
        elif text == "Questions & Feedback Welcome":
            set_box(shape, 0.8, 2.72, 11.73, 0.42)
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(18)
                    r.font.color.rgb = MUTED


def main():
    prs = Presentation(SOURCE)
    style_headers(prs)
    fix_title_slide(prs.slides[0])
    fix_outline(prs.slides[1])
    fix_slide_3(prs.slides[2])
    fix_experimental_setup(prs.slides[7])
    fix_gate_table(prs.slides[9])
    for number, slide in enumerate(prs.slides, 1):
        align_dense_cards(slide, number)
    fix_summary_cards(prs.slides[13])
    fix_end_slide(prs.slides[14])
    prs.save(OUTPUT)
    print(f"Saved {OUTPUT.resolve()}")


if __name__ == "__main__":
    main()
